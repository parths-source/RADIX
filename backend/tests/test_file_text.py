import io

import pytest
from docx import Document

from app.config import REPO_ROOT
from app.errors import AppError
from app.services.file_text import DOCX_MIME, extract_text, sanitize_filename, validate_upload

SAMPLES = REPO_ROOT / "data" / "samples"
MB = 1024 * 1024

# A valid one-page PDF with no text layer (stands in for a scanned PDF).
BLANK_PDF = (
    b"%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
    b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
    b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]>>endobj\n"
    b"trailer<</Root 1 0 R>>\n%%EOF\n"
)


def _code(exc_info):
    return exc_info.value.status_code, exc_info.value.code


@pytest.mark.parametrize(
    "raw, clean",
    [("../../etc/passwd", "passwd"), ("C:\\Users\\me\\My CV (final).pdf", "My_CV_final.pdf"),
     ("résumé.DOCX", "r_sum.docx"), ("...", "file")],
)
def test_sanitize_filename(raw, clean):
    assert sanitize_filename(raw) == clean


def test_accepts_docx_and_pdf():
    docx = (SAMPLES / "jds" / "google_software_engineer.docx").read_bytes()
    pdf = (SAMPLES / "jds" / "google_software_engineer.pdf").read_bytes()
    assert validate_upload("a.docx", DOCX_MIME, docx, 5 * MB) == ".docx"
    assert validate_upload("a.pdf", "application/pdf", pdf, 5 * MB) == ".pdf"


@pytest.mark.parametrize(
    "name, content_type, data, expected",
    [
        ("cv.txt", "text/plain", b"hello", (422, "unsupported_file_type")),
        ("cv.doc", "application/msword", b"\xd0\xcf", (422, "unsupported_file_type")),
        ("cv.pdf", "text/plain", b"%PDF-1.4", (422, "unsupported_file_type")),  # extension and type disagree
        ("cv.pdf", "application/pdf", b"not a pdf", (422, "unsupported_file_type")),  # bytes are not a PDF
        ("cv.pdf", "application/pdf", b"%PDF" + b"0" * (5 * MB), (413, "file_too_large")),
    ],
)
def test_rejects_bad_uploads(name, content_type, data, expected):
    with pytest.raises(AppError) as exc_info:
        validate_upload(name, content_type, data, 5 * MB)
    assert _code(exc_info) == expected


@pytest.mark.parametrize("folder", ["jds", "resumes"])
@pytest.mark.parametrize("ext", [".docx", ".pdf"])
def test_extracts_text_from_every_sample(folder, ext):
    for path in sorted((SAMPLES / folder).glob(f"*{ext}")):
        text = extract_text(path.read_bytes(), ext)
        assert len(text) > 300, path.name


def test_docx_tables_are_included():
    doc = Document()
    doc.add_paragraph("Skills")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Python", "Kubernetes"
    buffer = io.BytesIO()
    doc.save(buffer)
    assert "Python | Kubernetes" in extract_text(buffer.getvalue(), ".docx")


def test_pdf_without_text_is_rejected():
    with pytest.raises(AppError) as exc_info:
        extract_text(BLANK_PDF, ".pdf")
    assert _code(exc_info) == (422, "no_extractable_text")
    assert "Scanned PDFs" in exc_info.value.message


def test_corrupt_file_is_rejected():
    with pytest.raises(AppError) as exc_info:
        extract_text(b"PK\x03\x04 broken zip", ".docx")
    assert _code(exc_info) == (422, "unreadable_file")
